import json
from pathlib import Path
import re
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from lxml import etree

from app.config import Settings
from app.models.document import KnowledgeBaseDocument
from app.rag.chunking import chunk_documents
from app.rag.document_loader import DocumentExtractionError, load_knowledge_base_documents
from scripts import ingest_documents


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.document = Document()
        self.document.add_paragraph("Policy introduction.")

    def extract(self):
        self.document.save(self.root / "policy.docx")
        return load_knowledge_base_documents(self.root)[0]

    def define_list(self, identifier, levels):
        numbering = self.document.part.numbering_part.element
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), identifier)
        for index, (format_name, template, start) in enumerate(levels):
            level = OxmlElement("w:lvl")
            level.set(qn("w:ilvl"), str(index))
            for name, value in (("start", str(start)), ("numFmt", format_name), ("lvlText", template)):
                prop = OxmlElement("w:" + name)
                prop.set(qn("w:val"), value)
                level.append(prop)
            abstract.append(level)
        numbering.append(abstract)
        number = OxmlElement("w:num")
        number.set(qn("w:numId"), identifier)
        abstract_id = OxmlElement("w:abstractNumId")
        abstract_id.set(qn("w:val"), identifier)
        number.append(abstract_id)
        numbering.append(number)
        return abstract, number

    def add_list_item(self, identifier, level, text):
        paragraph = self.document.add_paragraph(text)
        properties = OxmlElement("w:numPr")
        for name, value in (("numId", identifier), ("ilvl", str(level))):
            prop = OxmlElement("w:" + name)
            prop.set(qn("w:val"), value)
            properties.append(prop)
        paragraph._p.get_or_add_pPr().append(properties)

    def test_blank_cells_headerless_tables_and_split_runs(self):
        table = self.document.add_table(rows=2, cols=3)
        table.cell(0, 1).text = "Domestic"
        table.cell(0, 2).text = "International"
        table.cell(1, 0).text = "Amount paid"
        table.cell(1, 1).text = "USD 15"
        paragraph = table.cell(1, 2).paragraphs[0]
        for text in ("USD ", "7", "5"):
            paragraph.add_run(text)
        self.document.add_paragraph("Subject to ticket conditions.")
        result = self.extract()
        self.assertFalse(result.extraction_warnings)
        self.assertIn("Row 1: Column 1: [empty] | Column 2: Domestic | Column 3: International", result.text)
        self.assertIn("Row 2: Column 1: Amount paid | Column 2: USD 15 | Column 3: USD 75", result.text)
        self.assertLess(result.text.index("introduction"), result.text.index("Table:"))
        self.assertLess(result.text.index("USD 75"), result.text.index("Subject to"))
        self.assertNotIn("declared header", result.text)

    def test_horizontal_and_vertical_merges(self):
        table = self.document.add_table(rows=3, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Cabin allowances"
        table.cell(0, 2).text = "Notes"
        table.cell(1, 0).merge(table.cell(2, 0)).text = "Africa"
        table.cell(1, 1).text = "Business"
        table.cell(2, 1).text = "Economy"
        result = self.extract()
        self.assertFalse(result.extraction_warnings)
        self.assertIn("Columns 1-2: Cabin allowances | Column 3: Notes", result.text)
        self.assertIn("Row 3: Column 1: [merged from row 2] Africa | Column 2: Economy", result.text)
        self.assertEqual(result.text.count("Cabin allowances"), 1)

    def test_declared_multirow_headers_are_not_inferred(self):
        table = self.document.add_table(rows=3, cols=2)
        for row in table.rows[:2]:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        table.cell(0, 0).text = "Cabin"
        table.cell(1, 0).text = "Class"
        table.cell(2, 0).text = "Economy"
        result = self.extract()
        self.assertEqual(result.text.count("(declared header)"), 2)
        self.assertIn("Row 3: Column 1: Economy", result.text)

    def test_nested_table_preserves_surrounding_cell_content(self):
        table = self.document.add_table(rows=1, cols=2)
        cell = table.cell(0, 0)
        cell.text = "Before nested table."
        nested = cell.add_table(rows=1, cols=2)
        nested.cell(0, 0).text = "Child category"
        nested.cell(0, 1).text = "10 kg"
        cell.add_paragraph("After nested table.")
        table.cell(0, 1).text = "Outer notes"
        result = self.extract()
        self.assertFalse(result.extraction_warnings)
        self.assertLess(result.text.index("Before"), result.text.index("Child category"))
        self.assertLess(result.text.index("10 kg"), result.text.index("After"))
        self.assertIn("Column 2: Outer notes", result.text)

    def test_omitted_grid_cells_keep_positions(self):
        table = self.document.add_table(rows=1, cols=3)
        row = table.rows[0]._tr
        row.remove(row.findall(qn("w:tc"))[0])
        row.remove(row.findall(qn("w:tc"))[-1])
        for name in ("gridBefore", "gridAfter"):
            prop = OxmlElement("w:" + name)
            prop.set(qn("w:val"), "1")
            row.get_or_add_trPr().append(prop)
        result = self.extract()
        self.assertIn("Column 1: [omitted] | Column 2: [empty] | Column 3: [omitted]", result.text)
        self.assertFalse(result.extraction_warnings)

    def test_headers_and_footers_are_extracted_once_as_supplementary_content(self):
        self.document.sections[0].header.paragraphs[0].text = "Policy version 2"
        self.document.sections[0].footer.paragraphs[0].text = "Contact support"
        self.document.add_section()
        result = self.extract()
        self.assertEqual(result.text.count("Policy version 2"), 1)
        self.assertEqual(result.text.count("Contact support"), 1)
        self.assertIn("Supplementary content", result.text)
        self.assertFalse(result.extraction_warnings)

    def test_footnotes_and_endnotes_keep_reference_markers(self):
        for kind, relationship in (("footnote", RT.FOOTNOTES), ("endnote", RT.ENDNOTES)):
            reference = OxmlElement(f"w:{kind}Reference")
            reference.set(qn("w:id"), "1")
            self.document.paragraphs[0].add_run()._r.append(reference)
            root = OxmlElement(f"w:{kind}s")
            note = OxmlElement(f"w:{kind}")
            note.set(qn("w:id"), "1")
            paragraph = OxmlElement("w:p")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = f"{kind} condition"
            run.append(text)
            paragraph.append(run)
            note.append(paragraph)
            root.append(note)
            part = Part(
                PackURI(f"/word/{kind}s.xml"),
                f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml",
                etree.tostring(root), self.document.part.package,
            )
            self.document.part.relate_to(part, relationship)
        result = self.extract()
        self.assertIn("[Footnote 1]", result.text)
        self.assertIn("Footnote 1:\nfootnote condition", result.text)
        self.assertIn("Endnote 1:\nendnote condition", result.text)
        self.assertFalse(result.extraction_warnings)

    def test_textbox_text_is_retained_but_requires_review(self):
        box = OxmlElement("w:txbxContent")
        paragraph = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Text box condition"
        run.append(text)
        paragraph.append(run)
        box.append(paragraph)
        self.document.paragraphs[0].add_run()._r.append(box)
        with self.assertLogs("app.rag.document_loader", level="WARNING"):
            result = self.extract()
        self.assertIn("Text box condition", result.text)
        self.assertTrue(result.extraction_warnings)
        with self.assertLogs("app.rag.chunking", level="WARNING"):
            self.assertEqual(chunk_documents([result], 1000, 150), [])

    def test_images_tracked_changes_fields_and_numbering_are_flagged(self):
        paragraph = self.document.paragraphs[0]
        for tag in ("drawing", "fldChar", "sym", "commentReference"):
            paragraph.add_run()._r.append(OxmlElement(f"w:{tag}"))
        insertion = OxmlElement("w:ins")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Unapproved condition"
        run.append(text)
        insertion.append(run)
        paragraph._p.append(insertion)
        paragraph._p.get_or_add_pPr().append(OxmlElement("w:numPr"))
        with self.assertLogs("app.rag.document_loader", level="WARNING"):
            result = self.extract()
        for expected in ("image", "field", "symbol", "comments", "Tracked changes", "list labels"):
            self.assertTrue(any(expected in warning for warning in result.extraction_warnings), expected)
        self.assertIn("Unapproved condition", result.text)

    def test_invalid_merges_and_table_grids_are_flagged(self):
        table = self.document.add_table(rows=1, cols=2)
        merge = OxmlElement("w:vMerge")
        table.cell(0, 0)._tc.get_or_add_tcPr().append(merge)
        table.rows[0]._tr.remove(table.rows[0]._tr.findall(qn("w:tc"))[-1])
        with self.assertLogs("app.rag.document_loader", level="WARNING"):
            result = self.extract()
        self.assertTrue(any("vertical merge" in warning for warning in result.extraction_warnings))
        self.assertTrue(any("Row width" in warning for warning in result.extraction_warnings))

    def test_legacy_merges_and_invalid_spans_require_review(self):
        table = self.document.add_table(rows=1, cols=1)
        properties = table.cell(0, 0)._tc.get_or_add_tcPr()
        properties.append(OxmlElement("w:hMerge"))
        span = OxmlElement("w:gridSpan")
        span.set(qn("w:val"), "invalid")
        properties.append(span)
        with self.assertLogs("app.rag.document_loader", level="WARNING"):
            result = self.extract()
        self.assertTrue(any("Legacy horizontal" in warning for warning in result.extraction_warnings))
        self.assertTrue(any("Invalid gridSpan" in warning for warning in result.extraction_warnings))

    def test_unmatched_note_and_hidden_text_require_review(self):
        run = self.document.paragraphs[0].add_run("Hidden condition")
        run.font.hidden = True
        reference = OxmlElement("w:footnoteReference")
        reference.set(qn("w:id"), "99")
        run._r.append(reference)
        with self.assertLogs("app.rag.document_loader", level="WARNING"):
            result = self.extract()
        self.assertIn("Hidden condition", result.text)
        self.assertTrue(any("Hidden text" in warning for warning in result.extraction_warnings))
        self.assertTrue(any("note" in warning and "missing" in warning for warning in result.extraction_warnings))

    def test_equations_and_embedded_blocks_require_review(self):
        equation = OxmlElement("m:oMath")
        text = OxmlElement("m:t")
        text.text = "x=1"
        equation.append(text)
        self.document.paragraphs[0]._p.append(equation)
        self.document.element.body.insert(1, OxmlElement("w:altChunk"))
        with self.assertLogs("app.rag.document_loader", level="WARNING"):
            result = self.extract()
        self.assertTrue(any("Equation" in warning for warning in result.extraction_warnings))
        self.assertTrue(any("Unsupported block" in warning for warning in result.extraction_warnings))

    def test_plain_numbered_lists_and_bullets_preserve_labels(self):
        numbering = self.document.part.numbering_part.element
        for format_name, identifier in (("decimal", "100"), ("bullet", "101")):
            abstract = OxmlElement("w:abstractNum")
            abstract.set(qn("w:abstractNumId"), identifier)
            level = OxmlElement("w:lvl")
            level.set(qn("w:ilvl"), "0")
            for name, value in (("start", "3"), ("numFmt", format_name), ("lvlText", "%1.")):
                prop = OxmlElement("w:" + name)
                prop.set(qn("w:val"), value)
                level.append(prop)
            abstract.append(level)
            numbering.append(abstract)
            number = OxmlElement("w:num")
            number.set(qn("w:numId"), identifier)
            abstract_id = OxmlElement("w:abstractNumId")
            abstract_id.set(qn("w:val"), identifier)
            number.append(abstract_id)
            numbering.append(number)
            for label in ("First", "Second"):
                paragraph = self.document.add_paragraph(f"{format_name} {label}")
                properties = paragraph._p.get_or_add_pPr()
                numpr = OxmlElement("w:numPr")
                num_id = OxmlElement("w:numId")
                num_id.set(qn("w:val"), identifier)
                numpr.append(num_id)
                properties.append(numpr)
        result = self.extract()
        self.assertFalse(result.extraction_warnings)
        self.assertIn("3. decimal First", result.text)
        self.assertIn("4. decimal Second", result.text)
        self.assertIn("[List level 1] - bullet First", result.text)

    def test_content_controls_are_visible_in_review_text_and_withheld(self):
        control = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        paragraph = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Controlled policy text"
        run.append(text)
        paragraph.append(run)
        content.append(paragraph)
        control.append(content)
        self.document.element.body.insert(1, control)
        with self.assertLogs("app.rag.document_loader", level="WARNING"):
            result = self.extract()
        self.assertIn("Controlled policy text", result.text)
        self.assertTrue(result.extraction_warnings)

    def test_letter_lists_preserve_case_punctuation_and_word_rollover(self):
        for identifier, fmt, template, labels in (
            ("100", "lowerLetter", "%1)", ("y)", "z)", "aa)", "bb)")),
            ("101", "upperLetter", "%1.", ("Y.", "Z.", "AA.", "BB.")),
        ):
            self.define_list(identifier, [(fmt, template, 25)])
            for index in range(4):
                self.add_list_item(identifier, 0, f"{fmt} item {index}")
        result = self.extract()
        self.assertFalse(result.extraction_warnings)
        for fmt, labels in (
            ("lowerLetter", ("y)", "z)", "aa)", "bb)")),
            ("upperLetter", ("Y.", "Z.", "AA.", "BB.")),
        ):
            for index, label in enumerate(labels):
                self.assertIn(f"{label} {fmt} item {index}", result.text)

    def test_nested_letter_lists_restart_after_parent_and_continue_after_plain_text(self):
        self.define_list("100", [("decimal", "%1.", 3), ("lowerLetter", "%2.", 1)])
        for level, text in ((0, "Parent one"), (1, "Child one"), (1, "Child two")):
            self.add_list_item("100", level, text)
        self.document.add_paragraph("A plain paragraph does not restart numbering.")
        for level, text in ((1, "Child three"), (0, "Parent two"), (1, "Restarted child")):
            self.add_list_item("100", level, text)
        result = self.extract()
        self.assertFalse(result.extraction_warnings)
        for expected in ("3. Parent one", "[List level 2] a. Child one",
                         "[List level 2] b. Child two", "[List level 2] c. Child three",
                         "4. Parent two", "[List level 2] a. Restarted child"):
            self.assertIn(expected, result.text)

    def test_nested_letters_restart_after_bullet_parent(self):
        self.define_list("100", [("bullet", "-", 1), ("lowerLetter", "%2)", 1)])
        for level, text in ((0, "Parent"), (1, "First"), (1, "Second"),
                            (0, "Next parent"), (1, "Restarted")):
            self.add_list_item("100", level, text)
        result = self.extract()
        self.assertFalse(result.extraction_warnings)
        self.assertIn("[List level 2] b) Second", result.text)
        self.assertIn("[List level 2] a) Restarted", result.text)

    def test_distinct_list_instances_have_independent_counters(self):
        abstract, _ = self.define_list("100", [("lowerLetter", "%1.", 1)])
        _, second = self.define_list("101", [("lowerLetter", "%1.", 1)])
        second.find(qn("w:abstractNumId")).set(qn("w:val"), abstract.get(qn("w:abstractNumId")))
        for identifier, text in (("100", "First list"), ("101", "Second list"), ("100", "Continued")):
            self.add_list_item(identifier, 0, text)
        result = self.extract()
        self.assertFalse(result.extraction_warnings)
        for expected in ("a. First list", "a. Second list", "b. Continued"):
            self.assertIn(expected, result.text)

    def test_unsupported_numbering_features_still_require_review(self):
        for identifier, feature in enumerate(("override", "restart", "compound", "roman", "legal", "zero"), 100):
            with self.subTest(feature=feature):
                key = str(identifier)
                fmt = "lowerRoman" if feature == "roman" else "lowerLetter"
                template = "%1.%2." if feature == "compound" else "%1."
                abstract, number = self.define_list(key, [(fmt, template, 0 if feature == "zero" else 1)])
                if feature == "override":
                    override = OxmlElement("w:lvlOverride")
                    override.set(qn("w:ilvl"), "0")
                    start = OxmlElement("w:startOverride")
                    start.set(qn("w:val"), "3")
                    override.append(start)
                    number.append(override)
                elif feature in {"restart", "legal"}:
                    prop = OxmlElement("w:lvlRestart" if feature == "restart" else "w:isLgl")
                    prop.set(qn("w:val"), "0" if feature == "restart" else "1")
                    abstract.find(qn("w:lvl")).append(prop)
                self.add_list_item(key, 0, f"{feature} item")
        with self.assertLogs("app.rag.document_loader", level="WARNING"):
            result = self.extract()
        self.assertEqual(len(result.extraction_warnings), 6)
        for feature in ("override", "restart", "compound", "roman", "legal", "zero"):
            self.assertIn(f"[List item] {feature} item", result.text)
        with self.assertLogs("app.rag.chunking", level="WARNING"):
            self.assertEqual(chunk_documents([result], 1000, 150), [])

    def test_empty_image_only_and_corrupt_documents_fail_explicitly(self):
        for variant in ("empty", "image", "corrupt"):
            with self.subTest(variant=variant):
                document = Document()
                if variant == "image":
                    document.add_paragraph().add_run()._r.append(OxmlElement("w:drawing"))
                path = self.root / "policy.docx"
                document.save(path)
                if variant == "corrupt":
                    path.write_bytes(b"not a Word package")
                with self.assertLogs("app.rag.document_loader", level="ERROR"):
                    with self.assertRaises(DocumentExtractionError):
                        load_knowledge_base_documents(self.root)

    def test_word_lock_files_are_skipped_and_duplicate_names_have_distinct_ids(self):
        for directory in ("one", "two"):
            folder = self.root / directory
            folder.mkdir()
            self.document.save(folder / "policy.docx")
        (self.root / "~$policy.docx").write_bytes(b"lock file")
        with self.assertLogs("app.rag.document_loader", level="WARNING"):
            results = load_knowledge_base_documents(self.root)
        self.assertEqual(len(results), 2)
        chunks = chunk_documents(results, 1000, 150)
        self.assertEqual(len({chunk.chunk_id for chunk in chunks}), 2)

    def test_ingestion_report_retains_review_text_but_excludes_flagged_chunks(self):
        clean = KnowledgeBaseDocument(source="clean.docx", file_type="docx", category="policy", text="Clean.")
        flagged = KnowledgeBaseDocument(
            source="flagged.docx", file_type="docx", category="policy",
            text="Review this content.", extraction_warnings=["Image needs review."],
        )
        settings = Settings(processed_data_dir=self.root)
        with patch.object(ingest_documents, "get_settings", return_value=settings):
            with patch.object(ingest_documents, "load_knowledge_base_documents", return_value=[clean, flagged]):
                with self.assertLogs("app.rag.chunking", level="WARNING"):
                    ingest_documents.main()
        report = json.loads((self.root / ingest_documents.OUTPUT_FILE_NAME).read_text(encoding="utf-8"))
        self.assertEqual(report["review_required_count"], 1)
        self.assertEqual(report["documents"][1]["extraction_status"], "needs_review")
        self.assertEqual(report["documents"][1]["text"], flagged.text)
        self.assertEqual({chunk["source"] for chunk in report["chunks"]}, {"clean.docx"})

    def test_empty_input_cannot_be_silently_chunked(self):
        empty = KnowledgeBaseDocument(source="empty.docx", file_type="docx", category="policy", text=" ")
        with self.assertLogs("app.rag.chunking", level="ERROR"):
            with self.assertRaises(ValueError):
                chunk_documents([empty], 1000, 150)


class KqDocumentRegressionTests(unittest.TestCase):
    def test_kq_letter_lists_are_preserved_and_all_documents_are_eligible(self):
        root = Path(__file__).resolve().parents[1] / "data" / "raw"
        documents = load_knowledge_base_documents(root)
        self.assertEqual(len(documents), 14)
        self.assertTrue(all(not doc.extraction_warnings for doc in documents))
        results = {Path(doc.source).name: doc.text for doc in documents}
        carriage = results["Conditions of Carriage.docx"]
        for expected in ("[List level 2] a. Our Customer Excellence Center",
                         "[List level 2] b. Kenya Airways website",
                         "[List level 2] c. At any local Kenya Airways office"):
            self.assertIn(expected, carriage)
        upgrade = results["Flight upgrade.docx"]
        for expected in ("a) The flight on which you were upgraded was cancelled or disrupted",
                         "b) The flight on which you were upgraded was cancelled and the airline",
                         "c) Your Offer was accepted and you were given an Upgrade",
                         "d) If the refund is approved"):
            self.assertIn(expected, upgrade)
        chunks = chunk_documents(documents, 1000, 150)
        self.assertEqual({chunk.source for chunk in chunks}, {doc.source for doc in documents})

    def test_editable_paragraph_content_is_retained_across_the_actual_knowledge_base(self):
        root = Path(__file__).resolve().parents[1] / "data" / "raw"
        with patch("app.rag.document_loader.logger"):
            documents = load_knowledge_base_documents(root)
        self.assertEqual(len(documents), 14)
        normalize = lambda value: re.sub(r"\s+", " ", value).strip()
        for document in documents:
            with self.subTest(source=document.source):
                output = normalize(document.text)
                with ZipFile(root / document.source) as package:
                    xml = etree.fromstring(package.read("word/document.xml"))
                # Compare each original editable paragraph, including table cells,
                # independently of the extractor's table formatting.
                for paragraph in xml.iter(qn("w:p")):
                    original = normalize("".join(
                        (node.text or "") if node.tag == qn("w:t") else " "
                        for node in paragraph.iter()
                        if node.tag in {qn("w:t"), qn("w:br"), qn("w:cr"), qn("w:tab")}
                    ))
                    if original:
                        self.assertTrue(original in output, f"Missing source paragraph in {document.source}: {original[:160]}")

    def test_kq_allowances_payment_columns_and_special_care_values(self):
        root = Path(__file__).resolve().parents[1] / "data" / "raw"
        with patch("app.rag.document_loader.logger"):
            results = {Path(doc.source).name: doc for doc in load_knowledge_base_documents(root)}
        baggage = results["Baggage Allowance Information.docx"].text
        self.assertIn("Column 1: Africa <-> America/Europe | Column 2: 2 PC 32kg/70lbs max each | Column 3: 2 PC 23kg/50lbs max each", baggage)
        payments = re.sub(r"\s+", " ", results["Payments.docx"].text)
        self.assertIn("Column 1: [empty] | Column 2: Domestic PNRs | Column 3: International PNRs", payments)
        self.assertIn("Column 1: Amount Paid | Column 2: USD 15 | Column 3: USD 75", payments)
        special = re.sub(r"\s+", " ", results["Special Care.docx"].text)
        self.assertIn("Column 1: Boeing 787 | Column 2: 152 x 139 x 149 cm", special)


if __name__ == "__main__":
    unittest.main()
