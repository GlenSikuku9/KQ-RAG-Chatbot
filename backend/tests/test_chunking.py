from difflib import SequenceMatcher
import json
import os
from pathlib import Path
import re
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from dotenv import dotenv_values
from pydantic import ValidationError

from app.config import Settings, get_settings
from app.models.document import DocumentBlock, KnowledgeBaseDocument
from app.rag.chunking import (
    CHUNKING_VERSION,
    _normalize_text,
    _split_spans,
    _split_with_overlap,
    chunk_documents,
)
from app.rag.document_loader import load_knowledge_base_documents
from scripts import ingest_documents


class ChunkingTests(unittest.TestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.document = Document()

    def load(self):
        self.document.save(self.root / "policy.docx")
        return load_knowledge_base_documents(self.root)[0]

    def test_invalid_settings_are_rejected_even_without_eligible_documents(self):
        for size, overlap in ((0, 0), (-1, 0), (10, -1), (10, 10), (10, 11),
                              (True, 0), (10, False), (10.5, 1), (10, 1.5)):
            with self.subTest(size=size, overlap=overlap):
                with self.assertLogs("app.config", level="ERROR"):
                    with self.assertRaises(ValueError):
                        chunk_documents([], size, overlap)
                with patch("app.config.logger"):
                    with self.assertRaises(ValidationError):
                        Settings(rag_chunk_size=size, rag_chunk_overlap=overlap)

    def test_environment_controls_size_and_overlap(self):
        self.addCleanup(get_settings.cache_clear)
        with patch.dict(os.environ, {"RAG_CHUNK_SIZE": "240", "RAG_CHUNK_OVERLAP": "35"}):
            get_settings.cache_clear()
            settings = get_settings()
            self.assertEqual((settings.rag_chunk_size, settings.rag_chunk_overlap), (240, 35))
        for size, overlap in (("no", "20"), ("100", "100"), ("100", "-1")):
            with patch.dict(os.environ, {"RAG_CHUNK_SIZE": size, "RAG_CHUNK_OVERLAP": overlap}):
                get_settings.cache_clear()
                with patch("app.config.logger"):
                    with self.assertRaises(ValueError):
                        get_settings()

    def test_default_settings_and_environment_example_agree(self):
        self.addCleanup(get_settings.cache_clear)
        with patch.dict(os.environ):
            os.environ.pop("RAG_CHUNK_SIZE", None)
            os.environ.pop("RAG_CHUNK_OVERLAP", None)
            get_settings.cache_clear()
            configured = get_settings()
        defaults = Settings()
        example = dotenv_values(Path(__file__).resolve().parents[1] / ".env.example")
        self.assertEqual((defaults.rag_chunk_size, defaults.rag_chunk_overlap), (1600, 250))
        self.assertEqual(configured.rag_chunk_size, defaults.rag_chunk_size)
        self.assertEqual(configured.rag_chunk_overlap, defaults.rag_chunk_overlap)
        self.assertEqual(example["RAG_CHUNK_SIZE"], str(defaults.rag_chunk_size))
        self.assertEqual(example["RAG_CHUNK_OVERLAP"], str(defaults.rag_chunk_overlap))

    def test_prose_keeps_all_text_and_whole_words_with_bounded_overlap(self):
        text = _normalize_text(
            "First paragraph introduces baggage rules and restrictions.\r\n\r\n"
            "Second paragraph describes ticket-specific allowances. Check the applicable route!\n\n"
            "Huduma kwa wateja inapatikana kila siku. Contact Kenya Airways for assistance."
        )
        for size, overlap in ((60, 0), (60, 15), (60, 59), (1000, 150)):
            with self.subTest(size=size, overlap=overlap):
                spans = _split_spans(text, size, overlap)
                chunks = _split_with_overlap(text, size, overlap)
                self.assertEqual(len(spans), len(chunks))
                covered = set()
                previous_end = 0
                for (start, end), chunk in zip(spans, chunks):
                    self.assertTrue(chunk)
                    self.assertLessEqual(len(chunk), size)
                    self.assertGreater(end, previous_end)
                    self.assertTrue(start == 0 or text[start - 1].isspace())
                    self.assertTrue(end == len(text) or text[end].isspace() or text[end - 1].isspace())
                    if previous_end:
                        self.assertLessEqual(max(0, previous_end - start), overlap)
                    covered.update(range(start, end))
                    previous_end = end
                self.assertTrue(all(index in covered for index, char in enumerate(text) if not char.isspace()))
                if overlap == 0:
                    self.assertEqual(" ".join(chunks).split(), text.split())

    def test_long_unbroken_tokens_are_not_cut_and_whitespace_is_normalized(self):
        token = "x" * 250
        chunks = _split_with_overlap("Start " + token + " end.", 40, 10)
        self.assertIn(token, chunks)
        self.assertEqual(_normalize_text(" \tA  B\r\n\r\n\r\nC  "), "A B\n\nC")
        self.assertEqual(_split_with_overlap("   ", 40, 10), [])

    def test_simple_table_chunks_repeat_opening_context_and_preserve_full_rows(self):
        self.document.add_paragraph("Fees depend on the itinerary.")
        table = self.document.add_table(rows=7, cols=3)
        table.cell(0, 1).text = "Domestic"
        table.cell(0, 2).text = "International"
        for row in range(1, 7):
            table.cell(row, 0).text = f"Fee {row}"
            table.cell(row, 1).text = f"USD {row * 10}"
            table.cell(row, 2).text = f"USD {row * 20}"
        self.document.add_paragraph("Subject to ticket conditions.")
        document = self.load()
        block = next(block for block in document.blocks if block.kind == "table")
        chunks = [chunk for chunk in chunk_documents([document], 400, 0) if chunk.table_location]
        self.assertGreater(len(chunks), 1)
        seen = []
        for chunk in chunks:
            self.assertIn("Opening row (context, not a declared header)", chunk.text)
            self.assertIn(block.rows[0].text, chunk.text)
            self.assertIn("Fees depend on the itinerary.", chunk.text)
            self.assertIn("Subject to ticket conditions.", chunk.text)
            self.assertEqual(chunk.table_location, block.source_location)
            for row in block.rows:
                if row.row_index in chunk.table_rows:
                    self.assertIn(row.text, chunk.text)
                    if row.row_index > 1:
                        seen.append(row.row_index)
        self.assertEqual(seen, list(range(2, 8)))

    def test_declared_multirow_headers_repeat_without_inference(self):
        table = self.document.add_table(rows=7, cols=2)
        for row in table.rows[:2]:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for index, row in enumerate(table.rows):
            row.cells[0].text = f"Category {index}"
            row.cells[1].text = f"Allowance {index}"
        document = self.load()
        chunks = chunk_documents([document], 290, 0)
        self.assertGreater(len(chunks), 1)
        for chunk in chunks:
            self.assertIn("Declared header rows:", chunk.text)
            self.assertNotIn("not a declared header", chunk.text)
            self.assertEqual(chunk.table_rows[:2], [1, 2])
            for row in document.blocks[0].rows[:2]:
                self.assertEqual(chunk.text.count(row.text), 1)
        self.assertEqual([row for chunk in chunks for row in chunk.table_rows if row > 2], list(range(3, 8)))

    def test_next_table_intro_is_not_repeated_in_the_previous_table(self):
        for title, amount in (("Name Correction", "USD 15"), ("Name Change", "USD 50")):
            self.document.add_paragraph(title)
            table = self.document.add_table(rows=2, cols=1)
            table.cell(0, 0).text = "Charge"
            table.cell(1, 0).text = amount
        document = self.load()
        chunks = [chunk for chunk in chunk_documents([document], 1000, 150) if chunk.table_location]
        self.assertEqual(len(chunks), 2)
        self.assertIn("Name Correction", chunks[0].text)
        self.assertNotIn("Name Change", chunks[0].text)
        self.assertIn("Name Change", chunks[1].text)

    def test_oversized_row_is_retained_without_withholding_the_document(self):
        table = self.document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Policy"
        table.cell(0, 1).text = "Condition"
        table.cell(1, 0).text = "Refund"
        table.cell(1, 1).text = "A complete policy condition. " * 60
        document = self.load()
        chunks = chunk_documents([document], 200, 30)
        self.assertFalse(document.extraction_warnings)
        self.assertEqual(len(chunks), 1)
        self.assertGreater(len(chunks[0].text), 200)
        self.assertIn(document.blocks[0].rows[1].text, chunks[0].text)

    def test_nested_and_merged_tables_remain_intact(self):
        table = self.document.add_table(rows=4, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Cabin allowances"
        table.cell(1, 0).merge(table.cell(2, 0)).text = "Africa"
        nested = table.cell(3, 0).add_table(rows=1, cols=2)
        nested.cell(0, 0).text = "Child category"
        nested.cell(0, 1).text = "10 kg"
        document = self.load()
        self.assertTrue(document.blocks[0].keep_together)
        chunks = chunk_documents([document], 100, 10)
        self.assertEqual(len(chunks), 1)
        self.assertIn(_normalize_text(document.blocks[0].text), chunks[0].text)
        self.assertEqual(chunks[0].table_rows, [1, 2, 3, 4])

    def test_table_row_overlap_uses_only_whole_rows_within_budget(self):
        table = self.document.add_table(rows=10, cols=1)
        for index, row in enumerate(table.rows):
            row.cells[0].text = f"Value {index}"
        document = self.load()
        chunks = chunk_documents([document], 235, 30)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(any(set(first.table_rows[1:]) & set(second.table_rows[1:])
                            for first, second in zip(chunks, chunks[1:])))
        for first, second in zip(chunks, chunks[1:]):
            overlap = set(first.table_rows[1:]) & set(second.table_rows[1:])
            repeated = [row.text for row in document.blocks[0].rows if row.row_index in overlap]
            self.assertLessEqual(len("\n".join(repeated)), 30)
            self.assertLessEqual(len(second.text), 235)
            self.assertTrue(set(second.table_rows) - set(first.table_rows))

    def test_headings_are_source_based_and_page_numbers_are_not_invented(self):
        self.document.add_heading("Baggage", 1)
        self.document.add_paragraph("Allowance information. " * 30)
        style = self.document.styles.add_style("Custom policy heading", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = self.document.styles["Heading 2"]
        self.document.add_paragraph("Exceptions", style=style)
        self.document.add_paragraph("Special conditions apply. " * 20)
        document = self.load()
        chunks = chunk_documents([document], 180, 20)
        self.assertEqual({chunk.section for chunk in chunks}, {"Baggage", "Baggage / Exceptions"})
        for chunk in chunks:
            self.assertIsNone(chunk.page)
            self.assertIn(f"Section: {chunk.section}", chunk.text)
            self.assertTrue(chunk.source_locations)
        self.assertTrue(all("Exceptions" not in chunk.text for chunk in chunks if chunk.section == "Baggage"))

    def test_supplementary_tables_remain_intact_and_identified(self):
        self.document.add_paragraph("Body policy.")
        table = self.document.sections[0].header.add_table(rows=4, cols=2, width=1000000)
        for row in table.rows:
            row.cells[0].text = "Header label"
            row.cells[1].text = "Header condition"
        document = self.load()
        chunks = chunk_documents([document], 80, 10)
        supplementary = [chunk for chunk in chunks if any("header" in location for location in chunk.source_locations)]
        self.assertEqual(len(supplementary), 1)
        self.assertIn("Supplementary content", supplementary[0].text)
        self.assertIn("Row 4:", supplementary[0].text)

    def test_long_headings_do_not_cause_tiny_repetitive_body_chunks(self):
        title = "Explicit heading " * 20
        self.document.add_heading(title, 1)
        self.document.add_paragraph("Short policy body.")
        chunks = chunk_documents([self.load()], 100, 10)
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].text.count(title.strip()), 1)
        self.assertIn("Short policy body.", chunks[0].text)

    def test_document_identity_versions_and_chunk_ids_are_deterministic(self):
        self.document.add_paragraph("Policy condition. " * 50)
        document = self.load()
        first = chunk_documents([document], 120, 15)
        repeated = chunk_documents([self.load()], 120, 15)
        self.assertEqual(first, repeated)
        self.assertEqual(len({chunk.chunk_id for chunk in first}), len(first))
        self.assertTrue(all(chunk.document_id == document.document_id and chunk.version == document.version
                            for chunk in first))
        changed = document.model_copy(update={
            "text": document.text + " Changed.", "blocks": [],
        })
        self.assertEqual(document.document_id, changed.document_id)
        self.assertNotEqual(document.version, changed.version)
        for alternative in (chunk_documents([changed], 120, 15), chunk_documents([document], 140, 15)):
            self.assertFalse({chunk.chunk_id for chunk in first} & {chunk.chunk_id for chunk in alternative})
        renamed = document.model_copy(update={"source": "another/policy.docx"})
        self.assertNotEqual(document.document_id, renamed.document_id)
        self.assertEqual(document.version, renamed.version)

    def test_plain_text_input_retains_metadata_without_guessing_sections(self):
        document = KnowledgeBaseDocument(
            source="folder/policy.docx", category="policy", file_type="docx",
            text="NOT AN EXPLICIT HEADING\n\nPolicy body.",
        )
        chunks = chunk_documents([document], 1000, 150)
        self.assertEqual(chunks[0].text, document.text)
        self.assertEqual(chunks[0].source, document.source)
        self.assertEqual(chunks[0].category, "policy")
        self.assertIsNone(chunks[0].section)
        self.assertIsNone(chunks[0].page)

    def test_duplicate_sources_and_inconsistent_blocks_fail_explicitly(self):
        document = KnowledgeBaseDocument(source="policy.docx", category="policy", file_type="docx", text="Policy.")
        with self.assertLogs("app.rag.chunking", level="ERROR"):
            with self.assertRaises(ValueError):
                chunk_documents([document, document], 100, 10)
        document.blocks = [DocumentBlock(kind="paragraph", text="Missing source content.", source_location="Body/block 1")]
        with self.assertLogs("app.rag.chunking", level="ERROR"):
            with self.assertRaises(ValueError):
                chunk_documents([document], 100, 10)

    def test_export_includes_versioned_document_and_chunk_metadata(self):
        self.document.add_paragraph("Policy body.")
        self.load()
        output = self.root / "processed"
        settings = Settings(raw_data_dir=self.root, processed_data_dir=output)
        with patch.object(ingest_documents, "get_settings", return_value=settings):
            ingest_documents.main()
        report = json.loads((output / ingest_documents.OUTPUT_FILE_NAME).read_text(encoding="utf-8"))
        self.assertEqual(report["chunking_version"], CHUNKING_VERSION)
        self.assertEqual(report["review_required_count"], 0)
        chunk = report["chunks"][0]
        document = report["documents"][0]
        self.assertEqual(chunk["document_id"], document["document_id"])
        self.assertEqual(chunk["version"], document["version"])
        self.assertTrue(chunk["source_locations"])
        self.assertIsNone(chunk["page"])


class KqChunkRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.documents = load_knowledge_base_documents(Path(__file__).resolve().parents[1] / "data" / "raw")
        cls.settings = Settings()
        cls.chunks = chunk_documents(cls.documents, cls.settings.rag_chunk_size, cls.settings.rag_chunk_overlap)

    def test_tuned_defaults_keep_each_corpus_paragraph_and_table_complete(self):
        for document in self.documents:
            chunks = [chunk for chunk in self.chunks if chunk.source == document.source]
            for block in document.blocks:
                with self.subTest(source=document.source, location=block.source_location):
                    if block.kind == "table":
                        related = [chunk for chunk in chunks if chunk.table_location == block.source_location]
                        self.assertEqual(len(related), 1)
                        self.assertEqual(related[0].table_rows, [row.row_index for row in block.rows])
                    else:
                        self.assertTrue(any(
                            block.source_location in chunk.source_locations
                            and _normalize_text(block.text) in chunk.text for chunk in chunks
                        ), "No chunk contains the complete source paragraph.")
        source_chars = sum(len(_normalize_text(document.text)) for document in self.documents)
        # Includes intentional overlap, table context, and source labels, not just duplication.
        self.assertLessEqual(sum(len(chunk.text) for chunk in self.chunks) / source_chars, 1.20)
        self.assertTrue(all(len(chunk.text) <= self.settings.rag_chunk_size for chunk in self.chunks))

    def test_all_actual_table_rows_keep_their_context_in_every_chunk(self):
        self.assertEqual(len(self.documents), 14)
        for document in self.documents:
            self.assertFalse(document.extraction_warnings)
            for block in document.blocks:
                if block.kind != "table":
                    continue
                with self.subTest(source=document.source, table=block.source_location):
                    chunks = [chunk for chunk in self.chunks
                              if chunk.source == document.source and chunk.table_location == block.source_location]
                    self.assertTrue(chunks)
                    self.assertEqual({row for chunk in chunks for row in chunk.table_rows},
                                     {row.row_index for row in block.rows})
                    for chunk in chunks:
                        for row in block.rows:
                            if row.row_index in chunk.table_rows:
                                self.assertIn(_normalize_text(row.text), chunk.text)
                        self.assertIn(_normalize_text(block.rows[0].text), chunk.text)
                    if block.keep_together:
                        self.assertEqual(len(chunks), 1)

    def test_payment_and_baggage_facts_remain_associated(self):
        payments = [chunk.text for chunk in self.chunks if Path(chunk.source).name == "Payments.docx"
                    and "Column 1: Amount Paid | Column 2: USD 15 | Column 3: USD 75" in chunk.text]
        self.assertTrue(payments)
        for text in payments:
            self.assertIn("Domestic PNRs", text)
            self.assertIn("International PNRs", text)
            self.assertIn("Column 1: Amount Paid | Column 2: USD 15 | Column 3: USD 75", re.sub(r"\s+", " ", text))
        baggage = [chunk.text for chunk in self.chunks
                   if "Column 1: Africa <-> America/Europe" in chunk.text]
        self.assertTrue(baggage)
        self.assertTrue(any("2 PC 32kg/70lbs" in text and "2 PC 23kg/50lbs" in text for text in baggage))

    def test_all_sources_are_traceable_and_repeated_runs_match(self):
        self.assertEqual({chunk.source for chunk in self.chunks}, {doc.source for doc in self.documents})
        self.assertEqual(len({chunk.chunk_id for chunk in self.chunks}), len(self.chunks))
        self.assertEqual(self.chunks, chunk_documents(
            self.documents, self.settings.rag_chunk_size, self.settings.rag_chunk_overlap,
        ))
        for chunk in self.chunks:
            self.assertTrue(chunk.text.strip())
            self.assertTrue(chunk.source_locations)
            self.assertIsNone(chunk.page)
            self.assertIsNone(chunk.section)

    def test_all_actual_prose_is_covered_without_dropped_text(self):
        for document in self.documents:
            chunks = [chunk for chunk in self.chunks if chunk.source == document.source]
            for block in document.blocks:
                if block.kind == "table":
                    continue
                with self.subTest(source=document.source, location=block.source_location):
                    original = _normalize_text(block.text)
                    related = [chunk.text for chunk in chunks if block.source_location in chunk.source_locations]
                    self.assertTrue(related)
                    if any(original in text for text in related):
                        continue
                    covered = set()
                    for text in related:
                        for match in SequenceMatcher(None, original, text, autojunk=False).get_matching_blocks():
                            covered.update(range(match.a, match.a + match.size))
                    self.assertTrue(all(index in covered for index, char in enumerate(original) if not char.isspace()))


if __name__ == "__main__":
    unittest.main()
