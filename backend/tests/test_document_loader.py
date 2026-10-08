from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document
from pydantic import ValidationError

from app.models.document import KnowledgeBaseDocument
from app.models.document_chunk import DocumentChunk
from app.rag.document_loader import (
    SUPPORTED_EXTENSIONS,
    _extract_text,
    load_knowledge_base_documents,
)


class DocxOnlyTests(unittest.TestCase):
    def setUp(self):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)

    def write_document(self, path: Path):
        document = Document()
        document.add_paragraph("Support policy.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Category"
        table.cell(0, 1).text = "Allowance"
        table.cell(1, 0).text = "Example"
        table.cell(1, 1).text = "One item"
        document.add_paragraph("Conditions apply.")
        document.save(path)

    def test_only_docx_is_supported(self):
        self.assertEqual(SUPPORTED_EXTENSIONS, {".docx"})

    def test_nested_uppercase_docx_preserves_existing_extraction(self):
        nested = self.root / "policies"
        nested.mkdir()
        self.write_document(nested / "Support-Policy.DOCX")
        documents = load_knowledge_base_documents(self.root)
        self.assertEqual(len(documents), 1)
        self.assertEqual(documents[0].file_type, "docx")
        self.assertEqual(documents[0].source, str(Path("policies") / "Support-Policy.DOCX"))
        self.assertEqual(documents[0].category, "Support Policy")
        self.assertEqual(
            documents[0].text,
            "Support policy.\n\nTable:\nRow 1: Column 1: Category | Column 2: Allowance"
            "\nRow 2: Column 1: Example | Column 2: One item"
            "\n\nConditions apply.",
        )

    def test_unsupported_direct_input_is_rejected_before_opening(self):
        with self.assertLogs("app.rag.document_loader", level="ERROR"):
            with self.assertRaisesRegex(ValueError, "Only DOCX"):
                _extract_text(self.root / "unsupported.bin")

    def test_mixed_batch_reports_skipped_files_and_loads_docx(self):
        self.write_document(self.root / "policy.docx")
        (self.root / "notes.txt").write_text("Not an accepted source.", encoding="utf-8")
        with self.assertLogs("app.rag.document_loader", level="WARNING") as logs:
            documents = load_knowledge_base_documents(self.root)
        self.assertEqual(len(documents), 1)
        self.assertIn("notes.txt", "\n".join(logs.output))
        self.assertIn("only DOCX", "\n".join(logs.output))

    def test_no_supported_documents_fails_explicitly(self):
        (self.root / "notes.txt").write_text("Unsupported.", encoding="utf-8")
        with self.assertLogs("app.rag.document_loader", level="WARNING"):
            with self.assertRaisesRegex(FileNotFoundError, r"\.docx"):
                load_knowledge_base_documents(self.root)

    def test_missing_and_empty_directories_remain_errors(self):
        for path in (self.root / "missing", self.root):
            with self.subTest(path=path):
                with self.assertRaises(FileNotFoundError):
                    load_knowledge_base_documents(path)

    def test_data_models_reject_other_formats(self):
        for model, payload in (
            (KnowledgeBaseDocument, {"source": "source", "category": "policy", "text": "text"}),
            (DocumentChunk, {
                "chunk_id": "source:0", "source": "source", "category": "policy",
                "document_id": "source-id", "version": "sha256:content",
                "chunk_index": 0, "text": "text",
            }),
        ):
            with self.subTest(model=model.__name__):
                self.assertEqual(model(**payload, file_type="docx").file_type, "docx")
                with self.assertRaises(ValidationError):
                    model(**payload, file_type="unsupported")


if __name__ == "__main__":
    unittest.main()
