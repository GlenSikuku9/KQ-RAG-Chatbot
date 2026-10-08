from copy import deepcopy
from datetime import datetime, timezone
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock, patch

from docx import Document
from docx.oxml import OxmlElement
from fastapi import HTTPException
from google.api_core.exceptions import Aborted, PermissionDenied
from google.cloud.firestore_v1 import Transaction

from app.config import Settings
from app.models.document import document_id_for_source
from app.rag.ingestion import (
    OUTPUT_FILE_NAME, REPORT_FILE_NAME, prepare_document,
    publish_document, remove_document, write_json_atomic,
)
from app.services.knowledge_base import KnowledgeBaseStore
from scripts import ingest_documents


class IngestionTests(unittest.TestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.raw = self.root / "raw"
        self.raw.mkdir()
        self.path = self.raw / "policy.docx"
        self.settings = Settings(raw_data_dir=self.raw, processed_data_dir=self.root / "processed")
        self.write_document("Policy conditions. " * 120)
        self.stored = {}
        self.references = {}
        self.transactions = []
        self.client = Mock()

        def reference(collection, identifier):
            key = f"{collection}/{identifier}"
            if key not in self.references:
                ref = Mock()
                ref.path = key
                ref.get.side_effect = lambda **kwargs: snapshot(ref, **kwargs)
                self.references[key] = ref
            return self.references[key]

        def collection(name):
            result = Mock()
            result.document.side_effect = lambda identifier: reference(name, identifier)
            def stream(**kwargs):
                output = []
                for key in sorted(self.stored):
                    if key.startswith(name + "/"):
                        item = snapshot(reference(name, key.split("/", 1)[1]), **kwargs)
                        item.id = key.split("/", 1)[1]
                        output.append(item)
                return output
            result.stream.side_effect = stream
            return result

        def snapshot(ref, transaction=None, **kwargs):
            if transaction is not None:
                self.assertFalse(transaction.pending, "Firestore reads must precede transaction writes.")
            result = Mock()
            result.exists = ref.path in self.stored
            result.to_dict.return_value = deepcopy(self.stored.get(ref.path))
            return result

        self.client.collection.side_effect = collection
        self.client.get_all.side_effect = lambda refs, **kwargs: [snapshot(ref, **kwargs) for ref in reversed(refs)]

        def transaction(**kwargs):
            tx = Mock(spec=Transaction)
            tx._id = b"test-ingestion"
            tx._read_only = False
            tx._max_attempts = kwargs["max_attempts"]
            tx.pending = []
            tx._begin.side_effect = lambda **kwargs: tx.pending.clear()
            tx.set.side_effect = lambda ref, payload: tx.pending.append(("set", ref.path, deepcopy(payload)))
            tx.delete.side_effect = lambda ref: tx.pending.append(("delete", ref.path, None))

            def commit():
                stored = deepcopy(self.stored)
                for action, key, payload in tx.pending:
                    if action == "set":
                        stored[key] = payload
                    else:
                        stored.pop(key, None)
                self.stored = stored
            tx._commit.side_effect = commit
            self.transactions.append(tx)
            return tx
        self.client.transaction.side_effect = transaction
        self.store = KnowledgeBaseStore(self.client)

    def write_document(self, text, warning=False):
        document = Document()
        document.add_paragraph(text)
        if warning:
            document.paragraphs[0].add_run()._r.append(OxmlElement("w:drawing"))
        document.save(self.path)

    def publish(self):
        return publish_document(self.path, self.settings, self.store)

    def chunks(self):
        return {key: value for key, value in self.stored.items() if key.startswith("document_chunks/")}

    def test_first_publish_records_metadata_and_chunks_with_native_utc_timestamps(self):
        result = self.publish()
        self.assertEqual(result.status, "processed")
        self.assertTrue(result.persisted)
        record = self.store.read("policy.docx")
        self.assertIsNotNone(record.active)
        self.assertEqual(record.active.chunk_count, len(self.chunks()))
        self.assertEqual(record.document_id, document_id_for_source("policy.docx"))
        self.assertEqual(record.created_at.tzinfo, timezone.utc)
        raw = self.stored[f"documents/{record.document_id}"]
        self.assertIsInstance(raw["created_at"], datetime)
        self.assertIsInstance(raw["active"]["published_at"], datetime)
        self.assertEqual(record.last_attempt.status, "processed")
        self.assertEqual(self.store.read_active_chunks("policy.docx"), prepare_document(self.path, self.settings).chunks)

    def test_index_snapshot_uses_only_active_versions_and_honors_removals(self):
        self.publish()
        initial = self.store.read_index_snapshot()
        self.assertTrue(initial)
        self.path.write_bytes(b"invalid")
        with self.assertLogs(level="ERROR"):
            self.publish()
        self.assertEqual(self.store.read_index_snapshot(), initial)
        remove_document("policy.docx", self.settings, self.store)
        self.assertEqual(self.store.read_index_snapshot(), [])

    def test_repeated_identical_input_performs_no_writes_or_duplicate_records(self):
        first = self.publish()
        stored = deepcopy(self.stored)
        second = self.publish()
        self.assertEqual(second.status, "unchanged")
        self.assertEqual(first.active_version, second.active_version)
        self.assertEqual(stored, self.stored)
        self.transactions[-1].set.assert_not_called()
        self.transactions[-1].delete.assert_not_called()

    def test_replacement_atomically_removes_old_chunks_and_preserves_creation_time(self):
        first = self.publish()
        before = self.store.read("policy.docx")
        old_ids = set(self.chunks())
        self.write_document("Replacement policy.")
        second = self.publish()
        self.assertEqual(second.status, "processed")
        self.assertNotEqual(first.active_version, second.active_version)
        self.assertFalse(old_ids & set(self.chunks()))
        self.assertEqual(self.store.read("policy.docx").created_at, before.created_at)
        self.assertEqual(len(self.chunks()), 1)
        self.assertEqual(self.store.read_active_chunks("policy.docx")[0].text, "Replacement policy.")

    def test_setting_change_replaces_chunks_even_when_document_version_is_unchanged(self):
        first = self.publish()
        old_ids = set(self.chunks())
        self.settings.rag_chunk_size = 700
        second = self.publish()
        self.assertEqual(second.active_version, first.active_version)
        self.assertEqual(second.status, "processed")
        self.assertFalse(old_ids & set(self.chunks()))
        self.assertEqual(self.store.read("policy.docx").active.chunking.size, 700)

    def test_failed_and_review_replacements_keep_last_successful_version(self):
        first = self.publish()
        stored_chunks = deepcopy(self.chunks())
        for review in (False, True):
            with self.subTest(review=review):
                if review:
                    self.write_document("New unreviewed policy.", warning=True)
                else:
                    self.path.write_bytes(b"not a Word package")
                with self.assertLogs(level="WARNING"):
                    result = self.publish()
                self.assertEqual(result.status, "needs_review" if review else "failed")
                self.assertTrue(result.persisted)
                self.assertEqual(result.active_version, first.active_version)
                self.assertEqual(self.chunks(), stored_chunks)
                record = self.store.read("policy.docx")
                self.assertEqual(record.last_attempt.status, result.status)
                self.assertEqual(record.active.version, first.active_version)
                self.assertEqual(len(self.store.read_active_chunks("policy.docx")), len(stored_chunks))

    def test_new_failed_document_has_no_active_chunks_and_can_recover(self):
        self.path.write_bytes(b"invalid")
        with self.assertLogs(level="ERROR"):
            failed = self.publish()
        self.assertTrue(failed.persisted)
        self.assertIsNone(failed.active_version)
        self.assertEqual(failed.active_chunk_count, 0)
        self.assertEqual(self.store.read_active_chunks("policy.docx"), [])
        self.write_document("Corrected policy.")
        self.assertEqual(self.publish().status, "processed")

    def test_reprocessing_previous_good_source_clears_failed_attempt(self):
        original = "Original policy."
        self.write_document(original)
        self.publish()
        self.write_document("Needs review", warning=True)
        with self.assertLogs(level="WARNING"):
            self.publish()
        previous_chunks = deepcopy(self.chunks())
        self.write_document(original)
        result = self.publish()
        self.assertEqual(result.status, "processed")
        self.assertEqual(self.chunks(), previous_chunks)
        self.assertEqual(self.store.read("policy.docx").last_attempt.status, "processed")

    def test_explicit_removal_is_idempotent_preserves_source_and_can_be_republished(self):
        self.publish()
        original = self.path.read_bytes()
        result = remove_document("policy.docx", self.settings, self.store)
        self.assertEqual(result.status, "removed")
        self.assertEqual(result.active_chunk_count, 0)
        self.assertFalse(self.chunks())
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(remove_document("policy.docx", self.settings, self.store).status, "unchanged")
        self.assertEqual(self.publish().status, "processed")

    def test_missing_source_is_not_implicitly_removed(self):
        self.publish()
        self.path.unlink()
        other = self.raw / "other.docx"
        document = Document()
        document.add_paragraph("Other policy.")
        document.save(other)
        publish_document(other, self.settings, self.store)
        self.assertIsNotNone(self.store.read("policy.docx").active)
        self.assertEqual(remove_document("policy.docx", self.settings, self.store).status, "removed")
        self.assertIsNotNone(self.store.read("other.docx").active)

    def test_unknown_removal_and_outside_root_paths_fail_explicitly(self):
        with self.assertLogs(level="WARNING"):
            result = remove_document("unknown.docx", self.settings, self.store)
        self.assertEqual(result.status, "failed")
        self.assertFalse(result.persisted)
        for path in (self.root / "outside.docx", self.raw / ".." / "outside.docx"):
            with self.assertLogs(level="ERROR"):
                with self.assertRaises(ValueError):
                    publish_document(path, self.settings, self.store)
        self.assertFalse(self.stored)

    def test_stale_preparation_cannot_overwrite_newer_publish_or_removal(self):
        self.publish()
        previous = self.store.read("policy.docx")
        self.write_document("Prepared but delayed.")
        delayed = prepare_document(self.path, self.settings)
        self.write_document("Newer committed policy.")
        self.publish()
        for remove in (False, True):
            with self.subTest(remove=remove):
                if remove:
                    remove_document("policy.docx", self.settings, self.store)
                stored = deepcopy(self.stored)
                with self.assertLogs(level="WARNING"):
                    with self.assertRaises(HTTPException) as caught:
                        self.store.save_attempt("policy.docx", previous.revision, delayed.attempt, delayed.chunks)
                self.assertEqual(caught.exception.status_code, 409)
                self.assertEqual(self.stored, stored)

    def test_concurrent_identical_first_publish_is_idempotent(self):
        prepared = prepare_document(self.path, self.settings)
        self.store.save_attempt("policy.docx", None, prepared.attempt, prepared.chunks)
        result = self.store.save_attempt("policy.docx", None, prepared.attempt, prepared.chunks)
        self.assertEqual(result.status, "unchanged")

    def test_stale_failure_cannot_overwrite_newer_success_metadata(self):
        self.publish()
        previous = self.store.read("policy.docx")
        self.path.write_bytes(b"invalid")
        with self.assertLogs(level="ERROR"):
            failed = prepare_document(self.path, self.settings)
        self.write_document("New successful policy.")
        self.publish()
        with self.assertLogs(level="WARNING"):
            with self.assertRaises(HTTPException):
                self.store.save_attempt("policy.docx", previous.revision, failed.attempt, [])
        self.assertEqual(self.store.read("policy.docx").last_attempt.status, "processed")

    def test_stale_removal_cannot_delete_newly_published_chunks(self):
        self.publish()
        previous = self.store.read("policy.docx")
        removal = previous.last_attempt.model_copy(update={"status": "removed", "version": None, "chunk_count": 0})
        self.write_document("Newer replacement.")
        self.publish()
        stored = deepcopy(self.stored)
        with self.assertLogs(level="WARNING"):
            with self.assertRaises(HTTPException) as caught:
                self.store.save_attempt("policy.docx", previous.revision, removal, [])
        self.assertEqual(caught.exception.status_code, 409)
        self.assertEqual(self.stored, stored)

    def test_commit_failure_is_explicit_and_old_version_is_untouched(self):
        self.publish()
        original = deepcopy(self.stored)
        self.write_document("Changed policy.")
        transaction_factory = self.client.transaction.side_effect
        def failing_transaction(**kwargs):
            tx = transaction_factory(**kwargs)
            tx._commit.side_effect = PermissionDenied("PRIVATE-DATA")
            return tx
        self.client.transaction.side_effect = failing_transaction
        with self.assertLogs(level="ERROR") as logs:
            result = self.publish()
        self.assertEqual(result.status, "failed")
        self.assertFalse(result.persisted)
        self.assertIsNone(result.active_chunk_count)
        self.assertEqual(self.stored, original)
        self.assertNotIn("PRIVATE-DATA", result.error + "\n".join(logs.output))

    def test_transaction_retry_does_not_duplicate_writes(self):
        factory = self.client.transaction.side_effect
        def retrying_transaction(**kwargs):
            tx = factory(**kwargs)
            commit = tx._commit.side_effect
            calls = 0
            def retry_commit():
                nonlocal calls
                calls += 1
                if calls == 1:
                    raise Aborted("conflict")
                commit()
            tx._commit.side_effect = retry_commit
            return tx
        self.client.transaction.side_effect = retrying_transaction
        result = self.publish()
        self.assertEqual(result.status, "processed")
        self.assertEqual(self.transactions[-1]._commit.call_count, 2)
        self.assertEqual(len(self.chunks()), result.active_chunk_count)

    def test_limits_record_failure_instead_of_partial_replacement(self):
        self.publish()
        original = deepcopy(self.chunks())
        self.write_document("Changed content. " * 300)
        with patch("app.services.knowledge_base.MAX_TRANSACTION_WRITES", 2):
            with self.assertLogs(level="ERROR"):
                result = self.publish()
        self.assertEqual(result.status, "failed")
        self.assertTrue(result.persisted)
        self.assertEqual(self.chunks(), original)
        self.assertIn("atomic storage limits", result.error)

    def test_record_and_transaction_byte_budgets_fail_without_truncation(self):
        self.write_document("Original short policy.")
        self.publish()
        original = deepcopy(self.chunks())
        self.write_document("x" * 10000)
        for limit in ("MAX_RECORD_BYTES", "MAX_TRANSACTION_BYTES"):
            with self.subTest(limit=limit):
                with patch("app.services.knowledge_base." + limit, 4000):
                    with self.assertLogs(level="ERROR"):
                        result = self.publish()
                self.assertEqual(result.status, "failed")
                self.assertTrue(result.persisted)
                self.assertEqual(self.chunks(), original)

    def test_cli_continues_publishing_good_documents_after_a_bad_one(self):
        self.path.write_bytes(b"invalid")
        good = Document()
        good.add_paragraph("Good policy.")
        good.save(self.raw / "valid.docx")
        with patch.object(ingest_documents, "get_settings", return_value=self.settings), \
                patch.object(ingest_documents, "get_firestore_client", return_value=self.client):
            with self.assertLogs(level="ERROR"):
                self.assertEqual(ingest_documents.main(["--publish"]), 1)
        report = json.loads((self.settings.processed_data_dir / REPORT_FILE_NAME).read_text())
        self.assertEqual(report["failed_count"], 1)
        self.assertEqual({result["status"] for result in report["results"]}, {"failed", "processed"})
        self.assertTrue(all(result["persisted"] for result in report["results"]))
        self.assertEqual(self.store.read("policy.docx").last_attempt.status, "failed")
        self.assertIsNotNone(self.store.read("valid.docx").active)

    def test_cli_review_required_exit_is_nonzero_even_with_retained_active_version(self):
        first = self.publish()
        self.write_document("Unreviewed replacement.", warning=True)
        with patch.object(ingest_documents, "get_settings", return_value=self.settings), \
                patch.object(ingest_documents, "get_firestore_client", return_value=self.client):
            with self.assertLogs(level="WARNING"):
                self.assertEqual(ingest_documents.main(["--publish"]), 1)
        self.assertEqual(self.store.read("policy.docx").active.version, first.active_version)

    def test_corrupt_record_or_active_chunks_are_not_silently_accepted(self):
        self.publish()
        record = self.store.read("policy.docx")
        saved = deepcopy(self.stored)
        key = f"documents/{record.document_id}"
        self.stored[key]["document_id"] = "wrong"
        with self.assertLogs(level="ERROR"):
            self.assertFalse(self.publish().persisted)
        self.stored = saved
        self.stored.pop(next(iter(self.chunks())))
        with self.assertLogs(level="ERROR"):
            with self.assertRaises(HTTPException) as caught:
                self.store.read_active_chunks("policy.docx")
        self.assertEqual(caught.exception.status_code, 500)
        with self.assertLogs(level="ERROR"):
            self.assertFalse(self.publish().persisted)

    def test_preview_continues_after_bad_document_and_never_accesses_firestore(self):
        self.path.write_bytes(b"invalid")
        good = Document()
        good.add_paragraph("Other good policy.")
        good.save(self.raw / "good.docx")
        with patch.object(ingest_documents, "get_settings", return_value=self.settings), \
                patch.object(ingest_documents, "get_firestore_client") as get_client:
            with self.assertLogs(level="ERROR"):
                self.assertEqual(ingest_documents.main([]), 1)
        get_client.assert_not_called()
        report = json.loads((self.settings.processed_data_dir / OUTPUT_FILE_NAME).read_text())
        self.assertEqual(report["failed_count"], 1)
        self.assertEqual(report["attempted_document_count"], 2)
        self.assertEqual(report["document_count"], 1)
        failed = next(item for item in report["processing_results"] if item["status"] == "failed")
        self.assertEqual(failed["source"], "policy.docx")
        self.assertFalse(failed["persisted"])
        self.assertEqual(report["purpose"], "offline_review_only")

    def test_cli_publish_and_remove_write_separate_durable_report(self):
        with patch.object(ingest_documents, "get_settings", return_value=self.settings), \
                patch.object(ingest_documents, "get_firestore_client", return_value=self.client):
            self.assertEqual(ingest_documents.main(["--publish"]), 0)
            self.assertFalse((self.settings.processed_data_dir / OUTPUT_FILE_NAME).exists())
            report_path = self.settings.processed_data_dir / REPORT_FILE_NAME
            report = json.loads(report_path.read_text())
            self.assertEqual(report["results"][0]["status"], "processed")
            self.assertEqual(ingest_documents.main(["--publish"]), 0)
            self.assertEqual(json.loads(report_path.read_text())["results"][0]["status"], "unchanged")
            self.assertEqual(ingest_documents.main(["--remove", "policy.docx"]), 0)
            self.assertEqual(json.loads(report_path.read_text())["results"][0]["status"], "removed")
            self.assertTrue(self.path.exists())

    def test_cli_storage_outage_and_report_write_failure_are_not_success(self):
        with patch.object(ingest_documents, "get_settings", return_value=self.settings), \
                patch.object(ingest_documents, "get_firestore_client", side_effect=HTTPException(503, "Unavailable")):
            with self.assertLogs(level="ERROR"):
                self.assertEqual(ingest_documents.main(["--publish"]), 1)
        output = self.root / "report.json"
        output.write_text('{"previous": true}')
        with patch("app.rag.ingestion.os.replace", side_effect=PermissionError("blocked")):
            with self.assertLogs(level="ERROR"):
                with self.assertRaises(PermissionError):
                    write_json_atomic(output, {"new": True})
        self.assertEqual(json.loads(output.read_text()), {"previous": True})
        self.assertFalse(list(self.root.glob("*.tmp")))

    def test_cli_output_failure_after_commit_reports_failure_but_retry_is_safe(self):
        with patch.object(ingest_documents, "get_settings", return_value=self.settings), \
                patch.object(ingest_documents, "get_firestore_client", return_value=self.client), \
                patch.object(ingest_documents, "write_json_atomic", side_effect=PermissionError("blocked")):
            with self.assertLogs(level="ERROR"):
                self.assertEqual(ingest_documents.main(["--publish"]), 1)
        self.assertIsNotNone(self.store.read("policy.docx").active)
        self.assertEqual(self.publish().status, "unchanged")


if __name__ == "__main__":
    unittest.main()
